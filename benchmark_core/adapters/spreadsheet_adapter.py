from pathlib import Path
from openpyxl import load_workbook
from .base import ArtifactAdapter


class SpreadsheetAdapter(ArtifactAdapter):
    def __init__(self, path: str | Path, data_only: bool = False):
        self.path = Path(path)
        self.wb = load_workbook(self.path, data_only=data_only)

    def get(self, selector: str):
        # selector syntax: Sheet Name!A1
        if "!" not in selector:
            if selector.startswith("sheet:"):
                name = selector.split(":", 1)[1]
                if name not in self.wb.sheetnames:
                    raise KeyError(name)
                return name
            raise KeyError(f"Spreadsheet selector must be 'Sheet!A1' or 'sheet:Name': {selector}")
        sheet_name, cell = selector.rsplit("!", 1)
        return self.wb[sheet_name][cell].value

    def number_format(self, selector: str):
        sheet_name, cell = selector.rsplit("!", 1)
        return self.wb[sheet_name][cell].number_format
